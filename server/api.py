from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import (
    APIRouter,
    Depends,
    File,
    Header,
    Query,
    Request,
    Response,
    UploadFile,
    WebSocket,
    WebSocketDisconnect,
)
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook

from config import config, update_config
from database.knowledge_service import KnowledgeService
from database.product_sync import ProductSyncService
from server.auth import AuthService, SESSION_COOKIE, Session
from server.errors import APIError
from server.schemas import (
    AccountLoginRequest,
    AccountUpdate,
    CustomerKnowledgeCreate,
    CustomerKnowledgeUpdate,
    KeywordCreate,
    KeywordImport,
    KeywordUpdate,
    LoginRequest,
    PlatformStatusUpdate,
    ProductKnowledgeUpdate,
    ProductSyncRequest,
    SettingsUpdate,
)
from service.account_service import account_service
from service.keyword_service import keyword_service
from utils.llm_provider import (
    ProfileValidationError,
    build_llm_profile,
    capability_confirmation_for,
    profile_to_dict,
    provider_choices,
    requires_tool_trust_confirmation,
    resolve_tool_capability,
)
from utils.logger_loguru import get_logger


router = APIRouter(prefix="/api/v1")
logger = get_logger("WebAPI")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000


def _auth(request: Request) -> AuthService:
    return request.app.state.auth


def require_session(request: Request) -> Session:
    return _auth(request).get_session(request.cookies.get(SESSION_COOKIE))


def require_csrf(
    request: Request,
    session: Session = Depends(require_session),
    x_csrf_token: str | None = Header(default=None),
) -> Session:
    _auth(request).verify_csrf(session, x_csrf_token)
    return session


def _pagination(page: int, page_size: int, items: list[Any]) -> dict[str, Any]:
    start = (page - 1) * page_size
    return {
        "items": items[start:start + page_size],
        "page": page,
        "page_size": page_size,
        "total": len(items),
    }


def _safe_account(account: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": int(account["id"]),
        "shop_db_id": int(account["shop_db_id"]),
        "channel_name": account["channel_name"],
        "shop_id": account["shop_id"],
        "shop_name": account["shop_name"],
        "shop_logo": account.get("shop_logo"),
        "user_id": account["user_id"],
        "username": account["username"],
        "status": account.get("status"),
        "has_password": bool(account.get("password")),
        "has_cookies": bool(account.get("cookies")),
    }


def _iso(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.astimezone().isoformat()
    return str(value)


def _product_dict(item: Any) -> dict[str, Any]:
    return {
        "id": item.id,
        "shop_id": item.shop_id,
        "goods_id": item.goods_id,
        "goods_name": item.goods_name,
        "price": item.price,
        "price_min": item.price_min,
        "price_max": item.price_max,
        "sold_quantity": item.sold_quantity,
        "thumb_url": item.thumb_url,
        "specifications": item.specifications,
        "extracted_content": item.extracted_content,
        "created_at": _iso(item.created_at),
        "updated_at": _iso(item.updated_at),
        "last_extracted_at": _iso(item.last_extracted_at),
    }


def _customer_knowledge_dict(item: Any) -> dict[str, Any]:
    return {
        "id": item.id,
        "shop_id": item.shop_id,
        "title": item.title,
        "content": item.content,
        "tags": item.tags,
        "enabled": item.enabled,
        "created_at": _iso(item.created_at),
        "updated_at": _iso(item.updated_at),
    }


def _settings_payload() -> dict[str, Any]:
    model = config.get_model().model_dump(mode="json")
    llm = dict(model["llm"])
    api_key = llm.pop("api_key", "")
    llm.pop("capability_confirmation", None)
    llm.pop("tool_trust_confirmation", None)
    llm["api_key_configured"] = bool(api_key)
    llm["api_key_masked"] = "********" if api_key else ""
    payload = {
        "llm": llm,
        "business_hours": model["business_hours"],
        "prompt": model["prompt"],
        "providers": [
            {"value": value, "label": label}
            for value, label in provider_choices()
        ],
    }
    version_source = json.dumps(
        {**payload, "api_key_configured": bool(api_key)},
        sort_keys=True,
        ensure_ascii=True,
    )
    payload["version"] = hashlib.sha256(version_source.encode("utf-8")).hexdigest()
    return payload


@router.post("/auth/login")
async def login(payload: LoginRequest, request: Request, response: Response):
    client_key = request.client.host if request.client else "unknown"
    session = request.app.state.auth.authenticate(
        payload.username, payload.password, client_key
    )
    response.set_cookie(
        SESSION_COOKIE,
        session.session_id,
        httponly=True,
        secure=request.app.state.auth.secure_cookie,
        samesite="strict",
        max_age=8 * 60 * 60,
        path="/",
    )
    return {"username": session.username, "csrf_token": session.csrf_token}


@router.post("/auth/logout")
async def logout(
    request: Request,
    response: Response,
    _: Session = Depends(require_csrf),
):
    request.app.state.auth.logout(request.cookies.get(SESSION_COOKIE))
    response.delete_cookie(SESSION_COOKIE, path="/")
    return {"ok": True}


@router.get("/auth/me")
async def me(session: Session = Depends(require_session)):
    return {"username": session.username, "csrf_token": session.csrf_token}


@router.get("/shops")
async def list_shops(
    _: Session = Depends(require_session),
):
    items = await asyncio.to_thread(account_service.get_all_shops)
    return {"items": items, "total": len(items)}


@router.get("/accounts")
async def list_accounts(
    search: str = "",
    shop_id: int | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: Session = Depends(require_session),
):
    accounts = await asyncio.to_thread(account_service.get_all_accounts_with_details)
    items = [_safe_account(account) for account in accounts]
    if shop_id is not None:
        items = [item for item in items if item["shop_db_id"] == shop_id]
    if search:
        needle = search.casefold()
        items = [
            item for item in items
            if needle in item["username"].casefold()
            or needle in item["shop_name"].casefold()
            or needle in str(item["user_id"]).casefold()
        ]
    return _pagination(page, page_size, items)


@router.post("/account-login-jobs", status_code=202)
async def create_account_login_job(
    payload: AccountLoginRequest,
    request: Request,
    _: Session = Depends(require_csrf),
):
    async def runner(report):
        report({"phase": "opening_browser", "message": "正在打开登录页面"})
        result = await account_service.login(
            payload.username, payload.password, headless=payload.headless
        )
        if not result:
            raise RuntimeError("platform login failed")
        report({"phase": "saving", "message": "正在保存账号"})
        account = await asyncio.to_thread(account_service.add_account_from_login, result)
        return account

    record = await request.app.state.jobs.submit(
        kind="account_login",
        resource_key=f"account-login:{payload.username.casefold()}",
        runner=runner,
    )
    return record.public()


@router.post("/accounts/{account_id}/verify", status_code=202)
async def verify_account(
    account_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    account = await asyncio.to_thread(account_service.get_account_by_id, account_id)
    if not account:
        raise APIError("ACCOUNT_NOT_FOUND", "账号不存在", 404)

    async def runner(report):
        from Channel.pinduoduo.pdd_login import refresh_pdd_cookies

        report({"phase": "verifying", "message": "正在验证平台登录状态"})
        result = await refresh_pdd_cookies(
            account["username"],
            account.get("password"),
            profile_scope=f"{account['channel_name']}:{account['shop_id']}:{account['user_id']}",
        )
        if not result or not result.get("cookies"):
            raise RuntimeError("account verification failed")
        updated = await asyncio.to_thread(
            account_service.update_account_cookies,
            account["channel_name"], account["shop_id"], account["user_id"],
            result["cookies"],
        )
        if not updated:
            raise RuntimeError("cookie update failed")
        return {"account_id": account_id, "verified": True}

    record = await request.app.state.jobs.submit(
        kind="account_verify",
        resource_key=f"account:{account_id}:verify",
        runner=runner,
    )
    return record.public()


@router.patch("/accounts/{account_id}")
async def update_account(
    account_id: int,
    payload: AccountUpdate,
    request: Request,
    _: Session = Depends(require_csrf),
):
    account = await asyncio.to_thread(account_service.get_account_by_id, account_id)
    if not account:
        raise APIError("ACCOUNT_NOT_FOUND", "账号不存在", 404)
    if request.app.state.runtimes.is_active(account_id):
        raise APIError("ACCOUNT_RUNTIME_BUSY", "请先停止该账号的自动回复", 409)
    values = payload.model_dump(exclude_unset=True)
    if values.get("password") == "":
        values.pop("password")
    if not values:
        return _safe_account(account)
    success = await asyncio.to_thread(
        account_service.update_account_info,
        account["channel_name"], account["shop_id"], account["user_id"],
        values.get("username"), values.get("password"), None, None,
    )
    if not success:
        raise APIError("INTERNAL_ERROR", "账号更新失败", 500)
    updated = await asyncio.to_thread(account_service.get_account_by_id, account_id)
    return _safe_account(updated)


@router.delete("/accounts/{account_id}")
async def delete_account(
    account_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    account = await asyncio.to_thread(account_service.get_account_by_id, account_id)
    if not account:
        raise APIError("ACCOUNT_NOT_FOUND", "账号不存在", 404)
    if request.app.state.runtimes.is_active(account_id):
        raise APIError("ACCOUNT_RUNTIME_BUSY", "请先停止该账号的自动回复", 409)
    success = await asyncio.to_thread(
        account_service.delete_account,
        account["channel_name"], account["shop_id"], account["user_id"],
    )
    if not success:
        raise APIError("INTERNAL_ERROR", "账号删除失败", 500)
    return {"deleted": True, "account_id": account_id}


@router.put("/accounts/{account_id}/platform-status")
async def set_platform_status(
    account_id: int,
    payload: PlatformStatusUpdate,
    request: Request,
    _: Session = Depends(require_csrf),
):
    from Channel.pinduoduo.utils.API.Set_up_online import AccountMonitor

    account = await asyncio.to_thread(account_service.get_account_by_id, account_id)
    if not account:
        raise APIError("ACCOUNT_NOT_FOUND", "账号不存在", 404)
    if payload.status == 3 and (
        request.app.state.runtimes.is_active(account_id)
        or await asyncio.to_thread(
            account_service.get_auto_reply_enabled, account_id
        )
    ):
        await request.app.state.runtimes.disable(account_id)

    def set_status() -> bool:
        monitor = AccountMonitor(
            account.get("cookies"), shop_id=account["shop_id"],
            user_id=account["user_id"], channel_name=account["channel_name"],
        )
        return monitor.set_csstatus(payload.status)

    if not await asyncio.to_thread(set_status):
        raise APIError("PROVIDER_UNAVAILABLE", "平台状态设置失败，请稍后重试", 502)
    if not await asyncio.to_thread(
        account_service.update_account_status,
        account["channel_name"], account["shop_id"], account["user_id"],
        payload.status,
    ):
        raise APIError("INTERNAL_ERROR", "本地状态更新失败", 500)
    return {"account_id": account_id, "status": payload.status}


@router.get("/runtimes")
async def list_runtimes(
    request: Request,
    _: Session = Depends(require_session),
):
    items = await request.app.state.runtimes.snapshots()
    return {"items": items, "total": len(items)}


@router.get("/runtimes/{account_id}")
async def get_runtime(
    account_id: int,
    request: Request,
    _: Session = Depends(require_session),
):
    account = await asyncio.to_thread(account_service.get_account_by_id, account_id)
    if not account:
        raise APIError("ACCOUNT_NOT_FOUND", "账号不存在", 404)
    return request.app.state.runtimes.snapshot_one(account)


@router.post("/runtimes/{account_id}/start", status_code=202)
async def start_runtime(
    account_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    return await request.app.state.runtimes.enable(account_id)


@router.post("/runtimes/{account_id}/stop", status_code=202)
async def stop_runtime(
    account_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    return await request.app.state.runtimes.disable(account_id)


@router.post("/runtimes/actions/start-all", status_code=202)
async def start_all_runtimes(
    request: Request,
    _: Session = Depends(require_csrf),
):
    return await request.app.state.runtimes.start_all()


@router.post("/runtimes/actions/stop-all", status_code=202)
async def stop_all_runtimes(
    request: Request,
    _: Session = Depends(require_csrf),
):
    return await request.app.state.runtimes.stop_all()


def _keyword_by_id(keyword_id: int) -> dict[str, Any] | None:
    return next(
        (item for item in keyword_service.get_all_keywords() if item["id"] == keyword_id),
        None,
    )


@router.get("/keywords")
async def list_keywords(
    search: str = "",
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: Session = Depends(require_session),
):
    items = await asyncio.to_thread(keyword_service.get_all_keywords)
    if search:
        needle = search.casefold()
        items = [item for item in items if needle in item["keyword"].casefold()]
    return _pagination(page, page_size, items)


@router.post("/keywords", status_code=201)
async def add_keyword(
    payload: KeywordCreate,
    _: Session = Depends(require_csrf),
):
    if not await asyncio.to_thread(keyword_service.add_keyword, payload.keyword):
        raise APIError("KEYWORD_ALREADY_EXISTS", "关键词已存在", 409)
    return await asyncio.to_thread(keyword_service.get_keyword, payload.keyword) if hasattr(keyword_service, "get_keyword") else {"keyword": payload.keyword}


@router.patch("/keywords/{keyword_id}")
async def update_keyword(
    keyword_id: int,
    payload: KeywordUpdate,
    _: Session = Depends(require_csrf),
):
    current = await asyncio.to_thread(_keyword_by_id, keyword_id)
    if not current:
        raise APIError("KEYWORD_NOT_FOUND", "关键词不存在", 404)
    if not await asyncio.to_thread(
        keyword_service.update_keyword, current["keyword"], payload.keyword
    ):
        raise APIError("KEYWORD_ALREADY_EXISTS", "关键词已存在", 409)
    return {"id": keyword_id, "keyword": payload.keyword}


@router.delete("/keywords/{keyword_id}")
async def remove_keyword(
    keyword_id: int,
    _: Session = Depends(require_csrf),
):
    current = await asyncio.to_thread(_keyword_by_id, keyword_id)
    if not current:
        raise APIError("KEYWORD_NOT_FOUND", "关键词不存在", 404)
    await asyncio.to_thread(keyword_service.delete_keyword, current["keyword"])
    return {"deleted": True, "id": keyword_id}


@router.post("/keywords/import")
async def import_keywords(
    payload: KeywordImport,
    _: Session = Depends(require_csrf),
):
    values = list(dict.fromkeys(line.strip() for line in payload.content.splitlines() if line.strip()))
    if len(values) > MAX_IMPORT_ROWS:
        raise APIError("UPLOAD_INVALID", "导入内容不能超过 5000 行", 422)
    added = skipped = failed = 0
    for value in values:
        if len(value) > 100:
            failed += 1
        elif await asyncio.to_thread(keyword_service.add_keyword, value):
            added += 1
        else:
            skipped += 1
    return {"added": added, "skipped": skipped, "failed": failed}


@router.get("/shops/{shop_id}/product-knowledge")
async def list_product_knowledge(
    shop_id: int,
    request: Request,
    search: str = "",
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: Session = Depends(require_session),
):
    service: KnowledgeService = request.app.state.knowledge
    items = await asyncio.to_thread(service.list_products_by_shop, shop_id)
    serialized = [_product_dict(item) for item in items]
    if search:
        needle = search.casefold()
        serialized = [
            item for item in serialized
            if needle in item["goods_name"].casefold()
            or needle in str(item["goods_id"])
        ]
    return _pagination(page, page_size, serialized)


@router.get("/product-knowledge/{product_id}")
async def get_product_knowledge(
    product_id: int,
    request: Request,
    _: Session = Depends(require_session),
):
    item = await asyncio.to_thread(request.app.state.knowledge.get_product_by_id, product_id)
    if not item:
        raise APIError("KNOWLEDGE_NOT_FOUND", "商品知识不存在", 404)
    return _product_dict(item)


@router.patch("/product-knowledge/{product_id}")
async def update_product_knowledge(
    product_id: int,
    payload: ProductKnowledgeUpdate,
    request: Request,
    _: Session = Depends(require_csrf),
):
    item = await asyncio.to_thread(
        request.app.state.knowledge.update_product,
        product_id, payload.specifications, payload.extracted_content,
    )
    if not item:
        raise APIError("KNOWLEDGE_NOT_FOUND", "商品知识不存在", 404)
    return _product_dict(item)


@router.delete("/product-knowledge/{product_id}")
async def delete_product_knowledge(
    product_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    if not await asyncio.to_thread(request.app.state.knowledge.delete_product, product_id):
        raise APIError("KNOWLEDGE_NOT_FOUND", "商品知识不存在", 404)
    return {"deleted": True, "id": product_id}


@router.delete("/shops/{shop_id}/product-knowledge")
async def clear_product_knowledge(
    shop_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    count = await asyncio.to_thread(
        request.app.state.knowledge.clear_products_by_shop, shop_id
    )
    return {"deleted": count, "shop_id": shop_id}


@router.post("/shops/{shop_id}/product-sync-jobs", status_code=202)
async def create_product_sync_job(
    shop_id: int,
    payload: ProductSyncRequest,
    request: Request,
    _: Session = Depends(require_csrf),
):
    shop = await asyncio.to_thread(account_service.get_shop_by_id, shop_id)
    if not shop:
        raise APIError("SHOP_NOT_FOUND", "店铺不存在", 404)
    accounts = await asyncio.to_thread(account_service.get_all_accounts_with_details)
    account = next((item for item in accounts if item["shop_db_id"] == shop_id), None)
    if not account:
        raise APIError("ACCOUNT_NOT_FOUND", "该店铺没有可用账号", 409)
    sync_service = ProductSyncService(request.app.state.knowledge)

    async def runner(report):
        def on_progress(progress):
            report(asdict(progress))

        result = await sync_service.sync_shop(
            shop_id=int(shop["shop_id"]),
            shop_db_id=shop_id,
            user_id=account["user_id"],
            is_full_sync=payload.mode == "full",
            progress_callback=on_progress,
        )
        return asdict(result)

    record = await request.app.state.jobs.submit(
        kind="product_sync",
        resource_key=f"shop:{shop_id}:product-sync",
        runner=runner,
        cancel_callback=sync_service.cancel,
    )
    return record.public()


@router.get("/shops/{shop_id}/customer-knowledge")
async def list_customer_knowledge(
    shop_id: int,
    request: Request,
    search: str = "",
    tag: str = "",
    enabled: bool | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    _: Session = Depends(require_session),
):
    service: KnowledgeService = request.app.state.knowledge
    items = await asyncio.to_thread(service.list_customer_service_with_disabled, shop_id)
    serialized = [_customer_knowledge_dict(item) for item in items]
    if enabled is not None:
        serialized = [item for item in serialized if item["enabled"] is enabled]
    if tag:
        serialized = [item for item in serialized if tag.casefold() in (item["tags"] or "").casefold()]
    if search:
        needle = search.casefold()
        serialized = [
            item for item in serialized
            if needle in item["title"].casefold()
            or needle in item["content"].casefold()
        ]
    return _pagination(page, page_size, serialized)


@router.post("/shops/{shop_id}/customer-knowledge", status_code=201)
async def create_customer_knowledge(
    shop_id: int,
    payload: CustomerKnowledgeCreate,
    request: Request,
    _: Session = Depends(require_csrf),
):
    if not await asyncio.to_thread(account_service.get_shop_by_id, shop_id):
        raise APIError("SHOP_NOT_FOUND", "店铺不存在", 404)
    item = await asyncio.to_thread(
        request.app.state.knowledge.add_customer_service,
        shop_id, payload.title, payload.content, payload.tags, payload.enabled,
    )
    refreshed = await asyncio.to_thread(
        request.app.state.knowledge.get_customer_service_by_id, item.id
    )
    return _customer_knowledge_dict(refreshed)


@router.patch("/customer-knowledge/{knowledge_id}")
async def update_customer_knowledge(
    knowledge_id: int,
    payload: CustomerKnowledgeUpdate,
    request: Request,
    _: Session = Depends(require_csrf),
):
    values = payload.model_dump(exclude_unset=True)
    item = await asyncio.to_thread(
        request.app.state.knowledge.update_customer_service,
        knowledge_id,
        values.get("title"), values.get("content"), values.get("tags"),
        values.get("enabled"),
    )
    if not item:
        raise APIError("KNOWLEDGE_NOT_FOUND", "客服知识不存在", 404)
    return _customer_knowledge_dict(item)


@router.delete("/customer-knowledge/{knowledge_id}")
async def delete_customer_knowledge(
    knowledge_id: int,
    request: Request,
    _: Session = Depends(require_csrf),
):
    if not await asyncio.to_thread(
        request.app.state.knowledge.delete_customer_service, knowledge_id
    ):
        raise APIError("KNOWLEDGE_NOT_FOUND", "客服知识不存在", 404)
    return {"deleted": True, "id": knowledge_id}


def _parse_import_file(
    filename: str, content: bytes
) -> tuple[list[dict[str, Any]], int]:
    suffix = Path(filename).suffix.lower()
    rows: list[dict[str, Any]] = []
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        for row in csv.DictReader(io.StringIO(text)):
            rows.append(dict(row))
    elif suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values, ())]
        for values_row in values:
            rows.append(dict(zip(headers, values_row)))
        workbook.close()
    else:
        raise APIError("UPLOAD_INVALID", "只支持 CSV 或 XLSX 文件", 422)
    if len(rows) > MAX_IMPORT_ROWS:
        raise APIError("UPLOAD_INVALID", "导入文件不能超过 5000 行", 422)
    normalized = []
    invalid = 0
    for row in rows:
        title = str(row.get("title") or row.get("标题") or "").strip()
        body = str(row.get("content") or row.get("内容") or "").strip()
        tags = str(row.get("tags") or row.get("标签") or "").strip() or None
        if not title or not body or len(title) > 255 or len(body) > 500_000:
            invalid += 1
            continue
        normalized.append({"title": title, "content": body, "tags": tags})
    return normalized, invalid


@router.post("/shops/{shop_id}/customer-knowledge/import")
async def import_customer_knowledge(
    shop_id: int,
    request: Request,
    file: UploadFile = File(...),
    _: Session = Depends(require_csrf),
):
    if not await asyncio.to_thread(account_service.get_shop_by_id, shop_id):
        raise APIError("SHOP_NOT_FOUND", "店铺不存在", 404)
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise APIError("UPLOAD_INVALID", "上传文件不能超过 5 MB", 422)
    try:
        rows, invalid = await asyncio.to_thread(
            _parse_import_file, file.filename or "", content
        )
    except (UnicodeDecodeError, csv.Error, ValueError) as exc:
        raise APIError("UPLOAD_INVALID", "文件格式或内容无效", 422) from exc
    success, skipped = await asyncio.to_thread(
        request.app.state.knowledge.batch_import_customer_service, shop_id, rows
    )
    return {
        "added": success,
        "skipped": skipped,
        "invalid": invalid,
    }


@router.get("/settings")
async def get_settings(_: Session = Depends(require_session)):
    return _settings_payload()


@router.put("/settings")
async def save_settings(
    payload: SettingsUpdate,
    request: Request,
    _: Session = Depends(require_csrf),
):
    async with request.app.state.settings_lock:
        current = _settings_payload()
        if payload.version != current["version"]:
            raise APIError(
                "CONFIG_CONFLICT", "配置已被修改，请刷新后重新提交", 409
            )
        current_key = config.get("llm.api_key", "")
        llm_data = payload.llm.model_dump(exclude={"confirm_risk"})
        llm_data["api_key"] = payload.llm.api_key or current_key
        try:
            profile = build_llm_profile(llm_data, require_confirmation=False)
            if payload.llm.confirm_risk:
                fingerprint = capability_confirmation_for(profile)
                if resolve_tool_capability(profile).requires_confirmation:
                    profile = profile.__class__(
                        **{**profile.__dict__, "capability_confirmation": fingerprint}
                    )
                if requires_tool_trust_confirmation(profile):
                    profile = profile.__class__(
                        **{**profile.__dict__, "tool_trust_confirmation": fingerprint}
                    )
            profile = build_llm_profile(profile, require_confirmation=True)
        except ProfileValidationError as exc:
            raise APIError(
                "CONFIG_INVALID", exc.safe_message, 422,
                {"field": exc.field, "reason": exc.code},
            ) from exc
        update_config(
            {
                "llm": profile_to_dict(profile),
                "business_hours": payload.business_hours.model_dump(),
                "prompt": payload.prompt.model_dump(),
            },
            save=True,
        )
        return _settings_payload()


@router.get("/jobs/{job_id}")
async def get_job(
    job_id: str,
    request: Request,
    _: Session = Depends(require_session),
):
    return request.app.state.jobs.get(job_id).public()


@router.post("/jobs/{job_id}/cancel")
async def cancel_job(
    job_id: str,
    request: Request,
    _: Session = Depends(require_csrf),
):
    return (await request.app.state.jobs.cancel(job_id)).public()


@router.get("/logs")
async def get_logs(
    request: Request,
    level: str = "",
    module: str = "",
    search: str = "",
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=1, le=500),
    _: Session = Depends(require_session),
):
    items = request.app.state.log_buffer.query(
        level=level or None, module=module or None, search=search or None
    )
    items.reverse()
    return _pagination(page, page_size, items)


@router.get("/logs/export")
async def export_logs(
    request: Request,
    format: str = Query("json", pattern="^(json|csv|txt)$"),
    level: str = "",
    module: str = "",
    search: str = "",
    _: Session = Depends(require_session),
):
    items = request.app.state.log_buffer.query(
        level=level or None, module=module or None, search=search or None
    )
    if format == "json":
        content = json.dumps(items, ensure_ascii=False, indent=2)
        media_type = "application/json"
    elif format == "csv":
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=["timestamp", "level", "module", "message"])
        writer.writeheader()
        writer.writerows(items)
        content = output.getvalue()
        media_type = "text/csv"
    else:
        content = "\n".join(
            f"{item['timestamp']} | {item['level']} | {item['module']} | {item['message']}"
            for item in items
        )
        media_type = "text/plain"
    return StreamingResponse(
        iter([content.encode("utf-8-sig")]),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="logs.{format}"'},
    )


@router.delete("/logs/buffer")
async def clear_logs(
    request: Request,
    _: Session = Depends(require_csrf),
):
    request.app.state.log_buffer.clear()
    return {"cleared": True}


@router.websocket("/events")
async def events(websocket: WebSocket):
    try:
        session = websocket.app.state.auth.get_session(
            websocket.cookies.get(SESSION_COOKIE)
        )
    except APIError:
        await websocket.close(code=4401)
        return
    await websocket.accept()
    queue = await websocket.app.state.events.subscribe()
    try:
        await websocket.send_json(
            {
                "event_id": "connected",
                "type": "system.connected",
                "occurred_at": datetime.now().astimezone().isoformat(),
                "data": {"username": session.username},
            }
        )
        while True:
            try:
                event = await asyncio.wait_for(queue.get(), timeout=25)
            except asyncio.TimeoutError:
                event = {
                    "event_id": "heartbeat",
                    "type": "system.heartbeat",
                    "occurred_at": datetime.now().astimezone().isoformat(),
                    "data": {},
                }
            await websocket.send_json(event)
    except (WebSocketDisconnect, RuntimeError):
        pass
    finally:
        await websocket.app.state.events.unsubscribe(queue)


__all__ = ["router"]
